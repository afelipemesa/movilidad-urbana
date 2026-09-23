# -*- coding: utf-8 -*-
"""
TOP 10 DE ARTICULOS POR POSTURA EPISTEMICA

Complementa la Figura 2 (figura2_posturas.py). Sobre el mismo universo de
42.208 articulos (2006-2025, pertinencia >= 0,35), identifica para cada
postura (observar el desplazamiento, preguntar con categorias previas,
escuchar la experiencia) los articulos con mayor similitud semantica a ella:

  - los 10 mas cercanos de todo el corpus, con su orientacion predominante;
  - los 10 mas cercanos dentro de cada orientacion (tecnicista, ambiental,
    social-humana);
  - la composicion por orientacion de los 10, 100 y 1.000 mas cercanos.

Necesita los titulos y resumenes, asi que requiere dimensiones.xlsx (salida
del paso 1).

Entrada : dimensiones.xlsx
Salidas : outputs/tables/top10_articulos_por_postura.xlsx
"""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from cargar_corpus import aviso_sin_texto, cargar_corpus

UMBRAL = 0.35
TABLAS = Path("outputs/tables")
SALIDA = TABLAS / "top10_articulos_por_postura.xlsx"
DIMS = ["TECNICISTA", "AMBIENTAL", "SOCIAL_HUMANA"]
LAB = {"TECNICISTA": "Tecnicista", "AMBIENTAL": "Ambiental", "SOCIAL_HUMANA": "Social-humana"}
EJES = {"OBSERVACION_TERRENO": "Observar el desplazamiento",
        "INTERACCION_ESTRUCTURADA": "Preguntar con categorías previas",
        "INTERACCION_EXPERIENCIAL": "Escuchar la experiencia"}
c = lambda n: f"sim_{n}_avg"

# ------------------------------------------------------------------ datos
df, hay_texto = cargar_corpus()
if not hay_texto:
    aviso_sin_texto("Top 10 de articulos por postura")
    raise SystemExit(0)

for k in [c(x) for x in DIMS + list(EJES)] + ["sim_relevance_avg", "year_num"]:
    df[k] = pd.to_numeric(df[k], errors="coerce")
d = df[(df.year_num >= 2006) & (df.year_num <= 2025) & (df.sim_relevance_avg >= UMBRAL)].copy()
d["ORI"] = np.array(DIMS)[d[[c(x) for x in DIMS]].values.argmax(axis=1)]
d["Dimensión"] = d.ORI.map(LAB)


def fila(r, col, **extra):
    return {**extra, "Dimensión": r["Dimensión"], "Título": r["title"], "Año": int(r["year_num"]),
            "Similitud con la postura": round(r[col], 4), "Resumen": r["abstract"]}


resumen, top_global, top_dim = [], [], []
for eje, nombre in EJES.items():
    col = c(eje)
    orden = d.sort_values(col, ascending=False)
    for n in (10, 100, 1000):
        vc = orden.head(n)["Dimensión"].value_counts()
        resumen.append({"Postura": nombre, "Top N": n, **{v: int(vc.get(v, 0)) for v in LAB.values()}})
    for i, (_, r) in enumerate(orden.head(10).iterrows(), 1):
        top_global.append(fila(r, col, Postura=nombre, Rango=i))
    for dim in DIMS:
        for i, (_, r) in enumerate(orden[orden.ORI == dim].head(10).iterrows(), 1):
            top_dim.append(fila(r, col, Postura=nombre, Rango=i))

resumen = pd.DataFrame(resumen)
top_global = pd.DataFrame(top_global)
top_dim = pd.DataFrame(top_dim)[["Postura", "Dimensión", "Rango", "Título", "Año",
                                 "Similitud con la postura", "Resumen"]]

# ------------------------------------------------------------------ excel
TABLAS.mkdir(parents=True, exist_ok=True)
with pd.ExcelWriter(SALIDA, engine="openpyxl") as w:
    resumen.to_excel(w, sheet_name="Resumen", index=False)
    top_global.to_excel(w, sheet_name="Top10 por postura", index=False)
    top_dim.to_excel(w, sheet_name="Top10 postura x dimensión", index=False)
    for ws in w.book.worksheets:
        for cel in ws[1]:
            cel.font = Font(bold=True, color="FFFFFF")
            cel.fill = PatternFill("solid", fgColor="2F4F4F")
        ws.freeze_panes = "A2"
        for columna in ws.columns:
            h = columna[0].value
            ws.column_dimensions[columna[0].column_letter].width = {"Título": 70, "Resumen": 90, "Postura": 30}.get(h, 14)
            for cel in columna[1:]:
                cel.alignment = Alignment(wrap_text=h in ("Título", "Resumen"), vertical="top")
        if ws.title != "Resumen":
            for r in range(2, ws.max_row + 1):
                ws.row_dimensions[r].height = 60

print("composicion por orientacion de los articulos mas cercanos a cada postura:")
print(resumen.to_string(index=False))
print("\nnumero 1 de cada postura:")
for _, r in top_global[top_global.Rango == 1].iterrows():
    print(f"  {r.Postura}: [{r['Dimensión']}] {r['Título']} ({r['Año']})")
print("\n  guardada:", SALIDA)
